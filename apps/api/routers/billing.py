"""Platform-admin org management + manually-managed org licensing.

Payment/plan machinery (Razorpay, Plan, BillingPayment) has been removed —
see docs/razorpay-billing-removal.md. Access is now gated purely by a
per-org license the platform admin issues/renews/extends/suspends/
reactivates by hand (payment happens entirely off-platform). See
`deps.py`'s `enforce_org_license` for how an inactive license blocks a
whole org, and `workers/license_expiry_worker.py` for how `license_status`
auto-flips to "expired" once `license_expires_at` passes.
"""

from __future__ import annotations

import io
import re
from datetime import UTC, datetime, timedelta
from typing import Annotated
from uuid import UUID

import openpyxl
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import selectinload

from apps.api.channels.voice.org_numbers import replace_org_phone_numbers
from apps.api.channels.voice.plivo_provisioning import fire_and_forget_register_org_plivo_numbers
from apps.api.channels.whatsapp.webhook_registration import fire_and_forget_register_org_webhook
from apps.api.core.crypto import encrypt_secret
from apps.api.core.security import generate_login_token, hash_token
from apps.api.core.sessions import invalidate_user_sessions
from apps.api.db.models.account_user import AccountUser
from apps.api.db.models.org import Org
from apps.api.db.models.org_membership import OrgMembership
from apps.api.db.models.org_phone_number import OrgPhoneNumber
from apps.api.db.models.platform_settings import PlatformSettings
from apps.api.deps import (
    CurrentOrgDep,
    DbDep,
    RedisDep,
    _org_is_platform_admin_owned,
    verify_platform_admin,
)
from apps.api.schemas.billing import (
    ExtendLicenseIn,
    IssueLicenseIn,
    OrgAdminOut,
    OrgUpdateIn,
    PlatformSettingsOut,
    PlatformSettingsUpdateIn,
    ReactivateLicenseIn,
    RegenerateAdminTokenOut,
    RenewLicenseIn,
    SocialLinksOut,
    SuspendLicenseIn,
)
from apps.api.schemas.org_numbers import OrgPhoneNumberOut

# Platform-wide org directory + license administration — gated by
# X-Admin-Token or `AccountUser.is_superuser` specifically
# (verify_platform_admin), not the OR-guard admin.py uses, since that guard
# accepts ANY org's valid session, which would let any self-signed-up
# customer manage every other org's license.
PlatformAdminDep = Annotated[None, Depends(verify_platform_admin)]

router = APIRouter(prefix="/billing", tags=["billing"])


def _update_org_integrity_detail(exc: IntegrityError) -> str:
    message = str(exc.orig).lower()
    if (
        "uq_org_phone_numbers_provider_number" in message
        or "org_phone_numbers.provider" in message
        or "org_phone_numbers.phone_number" in message
    ):
        return "That number is already assigned to another organization"
    if "account_users.email" in message:
        return "That email is already used by another account"
    return "Could not update organization"


async def _load_org_admin_contact(
    db: DbDep, org_id: UUID
) -> tuple[str | None, str | None, str | None]:
    admin_result = await db.execute(
        select(AccountUser.email, AccountUser.full_name, AccountUser.mobile)
        .join(OrgMembership, OrgMembership.account_user_id == AccountUser.id)
        .where(OrgMembership.org_id == org_id, OrgMembership.role == "admin")
        .order_by(OrgMembership.created_at)
        .limit(1)
    )
    admin_row = admin_result.first()
    return admin_row if admin_row else (None, None, None)


def _as_aware_utc(value: datetime) -> datetime:
    """SQLite (used in tests) doesn't persist tzinfo on a DateTime(timezone=True)
    column, so a value round-tripped through it comes back naive even though
    every value ever written here is UTC — normalize before comparing against
    or serializing alongside another aware datetime. Postgres (prod) already
    returns aware values, so this is a no-op there."""
    return value if value.tzinfo is not None else value.replace(tzinfo=UTC)


def _org_admin_out(
    org: Org,
    *,
    seat_count: int,
    admin_email: str | None,
    admin_name: str | None,
    admin_mobile: str | None,
) -> OrgAdminOut:
    return OrgAdminOut(
        id=str(org.id),
        name=org.name,
        default_country_code=org.default_country_code,
        license_status=org.license_status,
        license_expires_at=_as_aware_utc(org.license_expires_at).isoformat() if org.license_expires_at else None,
        license_issued_at=_as_aware_utc(org.license_issued_at).isoformat() if org.license_issued_at else None,
        license_duration_days=org.license_duration_days,
        license_notes=org.license_notes,
        seat_count=seat_count,
        admin_email=admin_email,
        admin_name=admin_name,
        admin_mobile=admin_mobile,
        created_at=org.created_at.isoformat(),
        phone_numbers=[
            OrgPhoneNumberOut.model_validate(n, from_attributes=True) for n in org.phone_numbers
        ],
        enabled_features=org.enabled_features,
        max_team_members=org.max_team_members,
    )


@router.get("/orgs", response_model=list[OrgAdminOut])
async def list_orgs(db: DbDep, _admin: PlatformAdminDep) -> list[OrgAdminOut]:
    """Platform-wide org directory — every org on the platform, not just the
    caller's own (see PlatformAdminDep). A regular user never gets this view:
    every non-admin route is scoped to the caller's own org.

    Excludes the platform's own operating org (any org a superuser belongs
    to) — it isn't a customer org to manage/license/delete from this
    directory, it's the one running Veerox itself (see
    `_org_is_platform_admin_owned`).
    """
    platform_admin_org_ids = set(
        (
            await db.execute(
                select(OrgMembership.org_id)
                .join(AccountUser, AccountUser.id == OrgMembership.account_user_id)
                .where(AccountUser.is_superuser.is_(True))
            )
        )
        .scalars()
        .all()
    )

    stmt = select(Org).options(selectinload(Org.phone_numbers)).order_by(Org.created_at)
    if platform_admin_org_ids:
        stmt = stmt.where(Org.id.notin_(platform_admin_org_ids))
    result = await db.execute(stmt)
    orgs = result.scalars().all()

    out: list[OrgAdminOut] = []
    for org in orgs:
        seat_count_result = await db.execute(
            select(func.count()).select_from(OrgMembership).where(OrgMembership.org_id == org.id)
        )
        seat_count = seat_count_result.scalar_one()
        admin_email, admin_name, admin_mobile = await _load_org_admin_contact(db, org.id)
        out.append(
            _org_admin_out(
                org,
                seat_count=seat_count,
                admin_email=admin_email,
                admin_name=admin_name,
                admin_mobile=admin_mobile,
            )
        )
    return out


@router.patch("/orgs/{org_id}", response_model=OrgAdminOut)
async def update_org(
    org_id: UUID, payload: OrgUpdateIn, db: DbDep, _admin: PlatformAdminDep
) -> OrgAdminOut:
    """Platform-admin edit of an org's own profile — name and its dedicated
    calling/WhatsApp numbers (see OrgUpdateIn for why license fields aren't
    editable here). Backs the "Edit" action on the Organizations page.
    """
    if await _org_is_platform_admin_owned(db, org_id):
        raise HTTPException(status_code=403, detail="This organization cannot be edited")

    org_row = await db.get(Org, org_id)
    if org_row is None:
        raise HTTPException(status_code=404, detail="Organization not found")

    fields = payload.model_dump(exclude_unset=True)
    phone_numbers = fields.pop("phone_numbers", None)
    admin_fields = {
        key: fields.pop(key) for key in ("admin_email", "admin_name", "admin_mobile") if key in fields
    }
    for key in (
        "plivo_auth_id",
        "plivo_auth_token",
        "twilio_account_sid",
        "twilio_auth_token",
        "meta_app_id",
        "meta_app_secret",
        "meta_access_token",
        "meta_whatsapp_business_account_id",
        "meta_verify_token",
    ):
        fields.pop(key, None)

    plivo_creds_changed = False
    if payload.plivo_auth_id and payload.plivo_auth_token:
        org_row.plivo_auth_id = payload.plivo_auth_id.strip()
        org_row.plivo_auth_token_encrypted = encrypt_secret(payload.plivo_auth_token.strip())
        plivo_creds_changed = True
    if payload.twilio_account_sid and payload.twilio_auth_token:
        org_row.twilio_account_sid = payload.twilio_account_sid.strip()
        org_row.twilio_auth_token_encrypted = encrypt_secret(payload.twilio_auth_token.strip())

    meta_creds_changed = False
    if payload.meta_app_id and payload.meta_app_secret and payload.meta_access_token:
        org_row.meta_app_id = payload.meta_app_id.strip()
        org_row.meta_app_secret_encrypted = encrypt_secret(payload.meta_app_secret.strip())
        org_row.meta_access_token_encrypted = encrypt_secret(payload.meta_access_token.strip())
        if payload.meta_whatsapp_business_account_id:
            org_row.meta_whatsapp_business_account_id = payload.meta_whatsapp_business_account_id.strip()
        if payload.meta_verify_token:
            org_row.meta_verify_token_encrypted = encrypt_secret(payload.meta_verify_token.strip())
        meta_creds_changed = True

    if fields.get("default_country_code") is None:
        fields.pop("default_country_code", None)  # NOT NULL column — null means "leave as is"
    if "name" in fields:
        name = (fields["name"] or "").strip()
        if not name:
            raise HTTPException(status_code=400, detail="Organization name cannot be empty")
        fields["name"] = name
    for field, value in fields.items():
        setattr(org_row, field, value)
    if phone_numbers is not None:
        for entry in payload.phone_numbers or []:
            phone_number = re.sub(r"\D", "", entry.phone_number)
            number_owner_result = await db.execute(
                select(OrgPhoneNumber.id).where(
                    OrgPhoneNumber.provider == entry.provider,
                    OrgPhoneNumber.phone_number == phone_number,
                    OrgPhoneNumber.org_id != org_row.id,
                )
            )
            if number_owner_result.scalar_one_or_none() is not None:
                raise HTTPException(
                    status_code=409,
                    detail="That number is already assigned to another organization",
                )
        await replace_org_phone_numbers(db, org_row.id, payload.phone_numbers or [])

    if admin_fields:
        with db.no_autoflush:
            admin_user_result = await db.execute(
                select(AccountUser)
                .join(OrgMembership, OrgMembership.account_user_id == AccountUser.id)
                .where(OrgMembership.org_id == org_row.id, OrgMembership.role == "admin")
                .order_by(OrgMembership.created_at)
                .limit(1)
            )
            admin_user = admin_user_result.scalar_one_or_none()
            if admin_user is None:
                raise HTTPException(status_code=404, detail="This organization has no admin account to edit")
            if "admin_email" in admin_fields:
                admin_email = admin_fields["admin_email"].strip()
                email_owner_result = await db.execute(
                    select(AccountUser.id).where(
                        AccountUser.email == admin_email,
                        AccountUser.id != admin_user.id,
                    )
                )
                if email_owner_result.scalar_one_or_none() is not None:
                    raise HTTPException(status_code=409, detail="That email is already used by another account")
                admin_user.email = admin_email
            if "admin_name" in admin_fields:
                admin_user.full_name = (admin_fields["admin_name"] or "").strip() or None
            if "admin_mobile" in admin_fields:
                admin_user.mobile = (admin_fields["admin_mobile"] or "").strip() or None

    try:
        await db.commit()
    except IntegrityError as exc:
        await db.rollback()
        raise HTTPException(status_code=409, detail=_update_org_integrity_detail(exc)) from exc

    if phone_numbers is not None or plivo_creds_changed:
        fire_and_forget_register_org_plivo_numbers(org_row.id)
    if meta_creds_changed:
        fire_and_forget_register_org_webhook(org_row.id)

    result = await db.execute(
        select(Org).options(selectinload(Org.phone_numbers)).where(Org.id == org_row.id)
    )
    org_row = result.scalar_one()

    seat_count_result = await db.execute(
        select(func.count()).select_from(OrgMembership).where(OrgMembership.org_id == org_row.id)
    )
    seat_count = seat_count_result.scalar_one()
    admin_email, admin_name, admin_mobile = await _load_org_admin_contact(db, org_row.id)

    return _org_admin_out(
        org_row,
        seat_count=seat_count,
        admin_email=admin_email,
        admin_name=admin_name,
        admin_mobile=admin_mobile,
    )


_CALLING_NUMBERS_CELL_LIMIT = 5


def _format_calling_numbers(phone_numbers: list[OrgPhoneNumberOut]) -> str:
    """Render an org's dedicated Plivo/Twilio calling numbers for one Excel
    cell, capped so an org with a large pool (bulk-purchased campaign
    numbers can run into the dozens or more) doesn't produce one unreadable
    comma-joined wall of text. Excludes provider="whatsapp" rows — those get
    their own column (see _format_whatsapp_numbers) since they're a
    different kind of number (Meta's phone_number_id, not E.164). Each
    provider's default (the number outbound calls actually dial from — the
    only one that's operationally load-bearing day to day) is always shown;
    everything else fills the remaining slots up to the cap, with a "+N
    more" tail for whatever didn't fit. The full list is still available in
    the dashboard's Edit Org dialog.
    """
    calling_numbers = [n for n in phone_numbers if n.provider != "whatsapp"]
    if not calling_numbers:
        return ""
    defaults = [n for n in calling_numbers if n.is_default]
    rest = [n for n in calling_numbers if not n.is_default]
    ordered = defaults + rest
    shown = ordered[:_CALLING_NUMBERS_CELL_LIMIT]
    text = ", ".join(
        f"{n.provider}:{n.phone_number}{' (default)' if n.is_default else ''}" for n in shown
    )
    remaining = len(ordered) - len(shown)
    if remaining > 0:
        text += f", +{remaining} more"
    return text


def _format_whatsapp_numbers(phone_numbers: list[OrgPhoneNumberOut]) -> str:
    """Render an org's dedicated WhatsApp phone_number_id(s) for one Excel
    cell — usually just one or two, so no cap/truncation like
    _format_calling_numbers needs."""
    whatsapp_numbers = [n for n in phone_numbers if n.provider == "whatsapp"]
    if not whatsapp_numbers:
        return ""
    defaults = [n for n in whatsapp_numbers if n.is_default]
    rest = [n for n in whatsapp_numbers if not n.is_default]
    return ", ".join(
        f"{n.phone_number}{' (default)' if n.is_default else ''}" for n in defaults + rest
    )


@router.get("/orgs.xlsx")
async def export_orgs_xlsx(db: DbDep, _admin: PlatformAdminDep) -> StreamingResponse:
    """Same data as GET /billing/orgs, as a downloadable .xlsx workbook —
    backs the Organizations page's export button."""
    orgs = await list_orgs(db, _admin)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Organizations"
    sheet.append(
        [
            "organization",
            "admin_email",
            "license_status",
            "license_expires_at",
            "team_members",
            "created_at",
            "calling_numbers",
            "whatsapp_numbers",
        ]
    )
    # Default column width is ~8.4 chars — too narrow for the formatted
    # datetime cell below (Excel shows "####" rather than truncating a
    # numeric/date value that doesn't fit, unlike a text cell) and for a
    # typical org name/email address.
    for col, width in (
        ("A", 24), ("B", 28), ("C", 14), ("D", 18), ("E", 13), ("F", 18), ("G", 30), ("H", 22),
    ):
        sheet.column_dimensions[col].width = width
    for org in orgs:
        # org.created_at is OrgAdminOut's serialized .isoformat() string (see
        # list_orgs above) — parse it back to a real datetime so openpyxl
        # writes an actual date cell instead of left-aligned ISO text Excel
        # won't sort/filter as a date. Excel dates can't carry a timezone,
        # so tzinfo is dropped (org.created_at is UTC).
        created_at = datetime.fromisoformat(org.created_at).replace(tzinfo=None)
        license_expires_at = (
            datetime.fromisoformat(org.license_expires_at).replace(tzinfo=None)
            if org.license_expires_at
            else None
        )
        calling_numbers = _format_calling_numbers(org.phone_numbers)
        whatsapp_numbers = _format_whatsapp_numbers(org.phone_numbers)
        sheet.append(
            [
                org.name,
                org.admin_email or "",
                org.license_status,
                license_expires_at,
                org.seat_count,
                created_at,
                calling_numbers,
                whatsapp_numbers,
            ]
        )
        if license_expires_at is not None:
            sheet.cell(row=sheet.max_row, column=4).number_format = "yyyy-mm-dd hh:mm"
        sheet.cell(row=sheet.max_row, column=6).number_format = "yyyy-mm-dd hh:mm"

    buf = io.BytesIO()
    workbook.save(buf)
    buf.seek(0)

    stamp = datetime.now(UTC).strftime("%Y-%m-%d")
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="organizations-{stamp}.xlsx"'},
    )


@router.delete("/orgs/{org_id}", status_code=204)
async def delete_org(org_id: UUID, db: DbDep, redis: RedisDep, _admin: PlatformAdminDep) -> None:
    """Platform-admin hard-delete of any org on the platform.

    Every org_id-scoped table (Users, Leads, Conversations, Appointments,
    CallCampaigns/CampaignTargets, OrgMemberships, FollowUpTasks, etc.) is
    ON DELETE CASCADE at the DB level (see migrations/), so deleting this
    one row cleans up everything under it in a single transaction. Refuses
    to delete the platform's own operating org (see
    `_org_is_platform_admin_owned`) — it's already hidden from GET
    /billing/orgs, but this is the actual enforcement, not just the list
    filter, in case a stale org_id ever reaches here another way.

    Sessions for every member are invalidated after the commit — org_id
    lives in the session payload itself (not re-checked against the DB per
    request, see deps.py's `get_current_org`), so a stale token would
    otherwise keep resolving to a now-nonexistent org instead of failing
    cleanly.
    """
    if await _org_is_platform_admin_owned(db, org_id):
        raise HTTPException(status_code=403, detail="This organization cannot be deleted")

    org_row = await db.get(Org, org_id)
    if org_row is None:
        raise HTTPException(status_code=404, detail="Organization not found")

    member_ids = (
        (await db.execute(select(OrgMembership.account_user_id).where(OrgMembership.org_id == org_id)))
        .scalars()
        .all()
    )

    await db.delete(org_row)
    await db.commit()

    for account_user_id in member_ids:
        await invalidate_user_sessions(redis, account_user_id)


@router.post("/orgs/{org_id}/regenerate-admin-token", response_model=RegenerateAdminTokenOut)
async def regenerate_admin_token(
    org_id: UUID, db: DbDep, redis: RedisDep, _admin: PlatformAdminDep
) -> RegenerateAdminTokenOut:
    """For when an admin-issued login token is lost — there's no password
    reset to fall back on (see routers/auth.py), and the original token
    can't be recovered since only its hash is stored. Issues a fresh token
    for the org's earliest admin account and immediately invalidates their
    existing sessions, since the old token may be compromised too.
    """
    admin_result = await db.execute(
        select(AccountUser)
        .join(OrgMembership, OrgMembership.account_user_id == AccountUser.id)
        .where(OrgMembership.org_id == org_id, OrgMembership.role == "admin")
        .order_by(OrgMembership.created_at)
        .limit(1)
    )
    admin = admin_result.scalar_one_or_none()
    if admin is None:
        raise HTTPException(status_code=404, detail="Org has no admin account")

    login_token = generate_login_token()
    admin.token_hash = hash_token(login_token)
    await db.commit()
    await invalidate_user_sessions(redis, admin.id)

    return RegenerateAdminTokenOut(
        account_user_id=str(admin.id), email=admin.email, login_token=login_token
    )


# Fallback duration for a one-click renew/reactivate that doesn't specify
# `days` on an org that has never had a duration recorded
# (Org.license_duration_days is still null — e.g. it predates this field).
DEFAULT_LICENSE_DURATION_DAYS = 30


def _resolve_duration_days(org_row: Org, requested_days: int | None) -> int:
    if requested_days is not None:
        if requested_days <= 0:
            raise HTTPException(status_code=400, detail="days must be positive")
        return requested_days
    return org_row.license_duration_days or DEFAULT_LICENSE_DURATION_DAYS


async def _get_licensable_org(db: DbDep, org_id: UUID) -> Org:
    if await _org_is_platform_admin_owned(db, org_id):
        raise HTTPException(status_code=403, detail="This organization's license cannot be managed")
    org_row = await db.get(Org, org_id)
    if org_row is None:
        raise HTTPException(status_code=404, detail="Organization not found")
    return org_row


async def _license_out(db: DbDep, org_row: Org) -> OrgAdminOut:
    # org_row.phone_numbers is lazy-loaded — force it with a fresh eager-load
    # query rather than accessing the relationship directly, same as
    # list_orgs/update_org above (an unawaited lazy-load under AsyncSession
    # raises MissingGreenlet).
    result = await db.execute(
        select(Org).options(selectinload(Org.phone_numbers)).where(Org.id == org_row.id)
    )
    org_row = result.scalar_one()
    seat_count_result = await db.execute(
        select(func.count()).select_from(OrgMembership).where(OrgMembership.org_id == org_row.id)
    )
    seat_count = seat_count_result.scalar_one()
    admin_email, admin_name, admin_mobile = await _load_org_admin_contact(db, org_row.id)
    return _org_admin_out(
        org_row,
        seat_count=seat_count,
        admin_email=admin_email,
        admin_name=admin_name,
        admin_mobile=admin_mobile,
    )


@router.post("/orgs/{org_id}/license/issue", response_model=OrgAdminOut)
async def issue_license(
    org_id: UUID, payload: IssueLicenseIn, db: DbDep, _admin: PlatformAdminDep
) -> OrgAdminOut:
    """First-time (or from-scratch) license grant — sets the org active with
    a fresh expiry `days` from now. Also the right call to make for an org
    that predates licensing (backfilled with no expiry set)."""
    if payload.days <= 0:
        raise HTTPException(status_code=400, detail="days must be positive")
    org_row = await _get_licensable_org(db, org_id)
    org_row.license_status = "active"
    org_row.license_expires_at = datetime.now(UTC) + timedelta(days=payload.days)
    org_row.license_issued_at = datetime.now(UTC)
    org_row.license_duration_days = payload.days
    if payload.notes is not None:
        org_row.license_notes = payload.notes.strip() or None
    await db.commit()
    await db.refresh(org_row)
    return await _license_out(db, org_row)


@router.post("/orgs/{org_id}/license/renew", response_model=OrgAdminOut)
async def renew_license(
    org_id: UUID, payload: RenewLicenseIn, db: DbDep, _admin: PlatformAdminDep
) -> OrgAdminOut:
    """Set a new expiry `days` from now and bring the license back to
    active — the normal action once a client has paid for another period
    off-platform. `days` omitted (a one-click renew, see QuickRenewButton
    in the frontend) reuses whatever duration this org was last issued/
    renewed for."""
    org_row = await _get_licensable_org(db, org_id)
    duration_days = _resolve_duration_days(org_row, payload.days)
    org_row.license_status = "active"
    org_row.license_expires_at = datetime.now(UTC) + timedelta(days=duration_days)
    org_row.license_duration_days = duration_days
    await db.commit()
    await db.refresh(org_row)
    return await _license_out(db, org_row)


@router.post("/orgs/{org_id}/license/extend", response_model=OrgAdminOut)
async def extend_license(
    org_id: UUID, payload: ExtendLicenseIn, db: DbDep, _admin: PlatformAdminDep
) -> OrgAdminOut:
    """Add `days` on top of the org's current expiry (or from now, if it has
    none or has already passed) — for a short top-up without having to
    compute a new absolute date."""
    if payload.days <= 0:
        raise HTTPException(status_code=400, detail="days must be positive")
    org_row = await _get_licensable_org(db, org_id)
    now = datetime.now(UTC)
    current_expiry = _as_aware_utc(org_row.license_expires_at) if org_row.license_expires_at else None
    base = current_expiry if current_expiry and current_expiry > now else now
    org_row.license_expires_at = base + timedelta(days=payload.days)
    org_row.license_status = "active"
    await db.commit()
    await db.refresh(org_row)
    return await _license_out(db, org_row)


@router.post("/orgs/{org_id}/license/suspend", response_model=OrgAdminOut)
async def suspend_license(
    org_id: UUID, payload: SuspendLicenseIn, db: DbDep, _admin: PlatformAdminDep
) -> OrgAdminOut:
    """Immediately lock the org out regardless of its expiry date — e.g. a
    payment dispute. Distinct from letting it expire naturally so the admin
    can tell the two apart later."""
    org_row = await _get_licensable_org(db, org_id)
    org_row.license_status = "suspended"
    if payload.notes is not None:
        org_row.license_notes = payload.notes.strip() or None
    await db.commit()
    await db.refresh(org_row)
    return await _license_out(db, org_row)


@router.post("/orgs/{org_id}/license/reactivate", response_model=OrgAdminOut)
async def reactivate_license(
    org_id: UUID, payload: ReactivateLicenseIn, db: DbDep, _admin: PlatformAdminDep
) -> OrgAdminOut:
    """Lift a manual suspension. If the current expiry has already passed,
    a fresh `days`-from-now expiry is set — reusing the org's last
    issued/renewed duration (or DEFAULT_LICENSE_DURATION_DAYS) when `days`
    is omitted, same fallback as renew_license — otherwise the background
    worker would just flip it straight back to "expired" on its next pass."""
    org_row = await _get_licensable_org(db, org_id)
    now = datetime.now(UTC)
    expiry_has_passed = (
        org_row.license_expires_at is not None and _as_aware_utc(org_row.license_expires_at) <= now
    )
    if payload.days is not None or expiry_has_passed:
        duration_days = _resolve_duration_days(org_row, payload.days)
        org_row.license_expires_at = now + timedelta(days=duration_days)
        org_row.license_duration_days = duration_days
    org_row.license_status = "active"
    await db.commit()
    await db.refresh(org_row)
    return await _license_out(db, org_row)


async def _get_platform_settings(db: DbDep) -> PlatformSettings:
    record = await db.get(PlatformSettings, 1)
    if record is None:
        # Defensive fallback in case the seed row from the migration is
        # somehow missing — creates it on first read instead of 500ing.
        record = PlatformSettings(id=1, help_desk_script=None, social_links={})
        db.add(record)
        await db.commit()
        await db.refresh(record)
    return record


@router.get("/platform-settings", response_model=PlatformSettingsOut)
async def get_platform_settings(db: DbDep, _admin: PlatformAdminDep) -> PlatformSettingsOut:
    """Platform-wide help-desk script + social links, for the admin editor
    on the Organizations page."""
    record = await _get_platform_settings(db)
    return PlatformSettingsOut(
        help_desk_script=record.help_desk_script, social_links=record.social_links
    )


@router.patch("/platform-settings", response_model=PlatformSettingsOut)
async def update_platform_settings(
    payload: PlatformSettingsUpdateIn, db: DbDep, _admin: PlatformAdminDep
) -> PlatformSettingsOut:
    record = await _get_platform_settings(db)
    fields = payload.model_dump(exclude_unset=True)
    if "help_desk_script" in fields:
        value = fields["help_desk_script"]
        record.help_desk_script = value.strip() if value and value.strip() else None
    if "social_links" in fields:
        record.social_links = {k: v.strip() for k, v in (fields["social_links"] or {}).items() if v and v.strip()}
    await db.commit()
    await db.refresh(record)
    return PlatformSettingsOut(
        help_desk_script=record.help_desk_script, social_links=record.social_links
    )


@router.get("/social-links", response_model=SocialLinksOut)
async def get_social_links(org: CurrentOrgDep, db: DbDep) -> SocialLinksOut:
    """Read-only social links for client dashboards — any authenticated org
    member (not just admins) can see these."""
    _ = org  # CurrentOrgDep only used to require *some* valid session
    record = await _get_platform_settings(db)
    return SocialLinksOut(social_links=record.social_links)
